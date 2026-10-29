import datetime
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Appointment, Visit
from accounts.tests.helpers import make_patient, make_professional

from .models import advance_balance, PatientAllowance, SlotOrder, BillingEntry, PackagePayment, PatientPackage, Plan, use_package_session
from .views import revenue_for


class BillingBase(TestCase):
    def setUp(self):
        self.user, self.profile = make_professional()
        self.client.force_login(self.user)
        self.patient = make_patient(self.profile)
        self.today = timezone.localdate()

    def entry(self, **extra):
        data = {'entry_date': self.today.isoformat(), 'patient_name': 'Walk In', 'service': 'Consult', 'rate': '500', 'payment_mode': 'cash', **extra}
        return self.client.post(reverse('billing:entry-create'), data)

    def package(self, **extra):
        data = {'patient': self.patient.pk, 'duration': '2_week', 'total_sessions': 3, 'total_amount': '3000', 'amount_paid_now': '1000', 'start_date': self.today.isoformat(), **extra}
        return self.client.post(reverse('billing:package-create'), data)


class BillingEntryTests(BillingBase):
    def test_pages_render(self):
        for params in ({}, {'date': '2026-01-01'}, {'date': 'junk'}, {'new': 1}):
            self.assertEqual(self.client.get(reverse('billing:billing'), params).status_code, 200)

    def test_create_edit_delete_and_totals(self):
        self.entry()
        self.entry(patient_name='Second', rate='1500.50', payment_mode='online')
        page = self.client.get(reverse('billing:billing'))
        self.assertEqual(page.context['total'], Decimal('2000.50'))
        self.assertEqual(dict(page.context['by_mode'])['Online'], Decimal('1500.50'))
        first = BillingEntry.objects.get(patient_name='Walk In')
        self.assertTrue(first.invoice_number.startswith('BILL-'))
        self.client.post(reverse('billing:entry-edit', args=[first.pk]), {'entry_date': self.today.isoformat(), 'patient_name': 'Walk In', 'service': 'Consult', 'rate': '700', 'payment_mode': 'card'})
        first.refresh_from_db()
        self.assertEqual((first.rate, first.payment_mode), (700, 'card'))
        self.client.post(reverse('billing:entry-delete', args=[first.pk]))
        self.assertEqual(BillingEntry.objects.count(), 1)

    def test_validation(self):
        self.entry(rate='-5')
        self.entry(patient_name='X')
        self.entry(rate='')
        self.assertEqual(BillingEntry.objects.count(), 0)

    def test_linked_patient_fills_in_blank_details(self):
        self.entry(patient=self.patient.pk, patient_name=self.patient.name)
        e = BillingEntry.objects.get()
        self.assertEqual((e.patient, e.sex, e.contact_number, e.age), (self.patient, 'F', '9841000000', self.patient.age))
        self.assertContains(self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]), {'tab': 'billing'}), e.invoice_number)

    def test_cannot_link_another_professionals_patient(self):
        _, other_profile = make_professional('other')
        theirs = make_patient(other_profile, 'Theirs')
        self.entry(patient=theirs.pk, patient_name='Theirs')
        self.assertEqual(BillingEntry.objects.count(), 0)

    def test_entries_are_private_and_need_login(self):
        self.entry()
        entry = BillingEntry.objects.get()
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertNotContains(other.get(reverse('billing:billing')), 'Walk In')
        self.assertEqual(other.post(reverse('billing:entry-delete', args=[entry.pk])).status_code, 404)
        self.assertEqual(other.post(reverse('billing:entry-edit', args=[entry.pk]), {}).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('billing:billing')).status_code, 302)


class PlanContentTests(TestCase):
    def test_four_plans_are_seeded_with_the_published_content(self):
        plans = {p.code: p for p in Plan.objects.all()}
        self.assertEqual(set(plans), {'client', 'professional', 'clinic', 'enterprise'})
        self.assertTrue(plans['client'].is_free)
        self.assertEqual((plans['professional'].price_monthly, plans['professional'].patient_limit), (0, None))  # free for the first 10 patients
        self.assertIn('200', plans['professional'].price_note)
        self.assertTrue(plans['clinic'].is_custom and plans['enterprise'].is_custom)

    def test_package_page_and_landing_show_the_database_plans(self):
        self.assertContains(self.client.get(reverse('accounts:package')), 'First 10 patients free')
        Plan.objects.filter(code='professional').update(price_monthly=Decimal('1500'))
        self.assertContains(self.client.get(reverse('accounts:package')), '<span class="amount">1,500</span>')
        self.assertContains(self.client.get('/'), '1500')


class PatientAllowanceTests(TestCase):
    """Professional pricing: 10 patients free, then NPR 200 per extra patient, bought as slots."""

    def setUp(self):
        self.user, self.profile = make_professional('allowed')
        self.client.force_login(self.user)
        self.new = lambda name='Extra': {'name': name, 'date_of_birth': '1980-05-05', 'gender': 'M', 'phone': '9841234567'}

    def fill(self, count):
        for i in range(count):
            make_patient(self.profile, f'Patient {i}')

    def test_defaults_are_ten_free_at_two_hundred(self):
        allowance = PatientAllowance.for_profile(self.profile)
        self.assertEqual((allowance.free_limit, allowance.paid_slots, allowance.rate, allowance.allowed, allowance.remaining), (10, 0, 200, 10, 10))

    def test_the_eleventh_patient_needs_a_slot(self):
        self.fill(9)
        self.assertEqual(self.client.post(reverse('accounts:patient-add'), self.new('Tenth')).status_code, 302)
        self.assertEqual(self.profile.patients.count(), 10)  # the 10th is still free
        r = self.client.post(reverse('accounts:patient-add'), self.new('Eleventh'))
        self.assertRedirects(r, reverse('accounts:settings') + '?tab=plan', fetch_redirect_response=False)
        self.assertEqual(self.profile.patients.count(), 10)
        self.assertRedirects(self.client.get(reverse('accounts:patient-add')), reverse('accounts:settings') + '?tab=plan', fetch_redirect_response=False)

    def test_buying_slots_unlocks_exactly_that_many(self):
        self.fill(10)
        self.client.post(reverse('accounts:settings'), {'tab': 'plan', 'quantity': '2'})
        order = SlotOrder.objects.get()
        self.assertEqual((order.quantity, order.amount, order.status), (2, 400, 'requested'))
        self.client.post(reverse('accounts:patient-add'), self.new('Waiting'))
        self.assertEqual(self.profile.patients.count(), 10)  # unpaid orders do not unlock anything
        self.assertTrue(order.mark_paid(self.user))
        self.assertFalse(order.mark_paid(self.user))  # cannot add the slots twice
        allowance = PatientAllowance.for_profile(self.profile)
        self.assertEqual((allowance.paid_slots, allowance.allowed), (2, 12))
        for name in ('Eleven', 'Twelve', 'Thirteen'):
            self.client.post(reverse('accounts:patient-add'), self.new(name))
        self.assertEqual(self.profile.patients.count(), 12)

    def test_bad_quantities_are_rejected(self):
        for bad in ('0', '-3', 'abc', '501', ''):
            self.client.post(reverse('accounts:settings'), {'tab': 'plan', 'quantity': bad})
        self.assertFalse(SlotOrder.objects.exists())

    def test_clinic_patients_are_not_counted_here(self):
        from clinics.models import Clinic, ClinicSubscription
        clinic = Clinic.objects.create(owner=self.profile, name='C', phone='015550001')
        ClinicSubscription.objects.create(professional=self.profile, status='active')
        self.fill(10)
        for i in range(3):
            make_patient(self.profile, f'Clinic {i}', clinic=clinic)
        self.assertEqual(PatientAllowance.for_profile(self.profile).used, 10)
        self.client.post(reverse('accounts:patient-add') + f'?clinic={clinic.pk}', {**self.new('At Clinic'), 'clinic': clinic.pk})
        self.assertEqual(self.profile.patients.filter(clinic=clinic).count(), 4)  # the clinic plan covers it

    def test_plan_tab_and_patient_page_show_usage(self):
        self.fill(8)
        page = self.client.get(reverse('accounts:settings'), {'tab': 'plan'})
        self.assertContains(page, '8</span> of 10 slots used')
        self.assertContains(page, 'NPR 200')
        self.assertContains(self.client.get(reverse('accounts:patients')), '8 of 10')  # near the limit: a heads-up on the list

    def test_legacy_professionals_with_many_patients_keep_working(self):
        self.fill(15)
        allowance = PatientAllowance.objects.create(professional=self.profile, paid_slots=5)  # what the migration records for them
        self.assertTrue(allowance.allowed == 15 and not allowance.can_add)

    def test_slot_orders_are_private_and_staff_mark_them_paid_in_the_admin(self):
        self.fill(1)
        other_user, other = make_professional('rival')
        order = SlotOrder.objects.create(professional=other, quantity=5, rate=200, amount=1000)
        self.assertNotContains(self.client.get(reverse('accounts:settings'), {'tab': 'plan'}), 'NPR 1000')
        staff = make_professional('staffer', staff=True)[0]
        staff.is_superuser = True
        staff.save()
        self.client.force_login(staff)
        r = self.client.post(reverse('admin:billing_slotorder_changelist'), {'action': 'mark_paid', '_selected_action': [order.pk]})
        order.refresh_from_db()
        self.assertEqual(order.status, 'paid')
        self.assertEqual(PatientAllowance.for_profile(other).paid_slots, 5)

    def test_the_numbers_come_from_settings(self):
        from django.test import override_settings
        with override_settings(FREE_PATIENT_LIMIT=3, PATIENT_SLOT_RATE=250):
            fresh_user, fresh = make_professional('fresh')
            allowance = PatientAllowance.for_profile(fresh)
            self.assertEqual((allowance.free_limit, allowance.rate), (3, 250))


class PackageTests(BillingBase):
    def test_all_four_durations_exist_with_the_right_length(self):
        self.assertEqual([label for _, label in PatientPackage.DURATION_CHOICES if '_week' in _], ['1 Week', '2 Weeks', '3 Weeks', '4 Weeks'])  # plus 5, 10 and 15 days and a typed number of days
        for key, days in (('1_week', 7), ('2_week', 14), ('3_week', 21), ('4_week', 28)):
            self.package(duration=key)
            pkg = PatientPackage.objects.latest('id')
            self.assertEqual(pkg.end_date - pkg.start_date, datetime.timedelta(days=days), key)

    def test_packages_page_renders_and_filters(self):
        self.package()
        for params in ({}, {'show': 'active'}, {'show': 'due'}, {'show': 'completed'}, {'show': 'bogus'}, {'new': 1}, {'patient': self.patient.pk}):
            self.assertEqual(self.client.get(reverse('billing:packages'), params).status_code, 200, params)
        page = self.client.get(reverse('billing:packages'))
        self.assertEqual((page.context['active_count'], page.context['sessions_left'], page.context['total_due'], page.context['collected']), (1, 3, 2000, 1000))
        self.assertEqual(len(self.client.get(reverse('billing:packages'), {'show': 'completed'}).context['packages']), 0)
        self.assertEqual(len(self.client.get(reverse('billing:packages'), {'show': 'due'}).context['packages']), 1)

    def test_start_package_with_partial_payment(self):
        self.package()
        pkg = PatientPackage.objects.get()
        self.assertEqual((pkg.amount_paid, pkg.amount_due, pkg.sessions_used, pkg.effective_status), (1000, 2000, 0, 'active'))
        self.assertEqual((pkg.payments.count(), pkg.professional), (1, self.profile))
        self.assertEqual(pkg.per_session, Decimal('1000'))

    def test_package_validation(self):
        self.package(amount_paid_now='5000')  # more than the price
        self.package(total_sessions=0)
        self.package(total_amount='-1')
        self.package(duration='5_week')
        self.package(patient='')
        self.assertEqual(PatientPackage.objects.count(), 0)

    def test_cannot_sell_a_package_to_someone_elses_patient(self):
        _, other_profile = make_professional('other')
        theirs = make_patient(other_profile, 'Theirs')
        self.package(patient=theirs.pk)
        self.assertEqual(PatientPackage.objects.count(), 0)

    def test_payments_are_capped_at_the_balance_and_stop_when_paid(self):
        self.package()
        pkg = PatientPackage.objects.get()
        url = reverse('billing:package-payment', args=[pkg.pk])
        self.client.post(url, {'amount': '99999', 'payment_mode': 'cash'})
        pkg.refresh_from_db()
        self.assertEqual((pkg.amount_paid, pkg.is_fully_paid), (3000, True))
        self.assertEqual(PackagePayment.objects.get(amount=2000).amount, 2000)  # capped, not 99,999
        self.client.post(url, {'amount': '10', 'payment_mode': 'cash'})
        self.assertEqual(pkg.payments.count(), 2)
        self.assertIsNone(pkg.record_payment(Decimal('-5')))

    def test_visits_use_sessions_until_the_package_completes(self):
        self.package()
        visit_url = reverse('accounts:patient-visit-add', args=[self.patient.pk])
        for _ in range(3):
            self.client.post(visit_url, {'date': self.today.isoformat(), 'session_type': 'treatment'})
        pkg = PatientPackage.objects.get()
        self.assertEqual((pkg.sessions_used, pkg.status, pkg.effective_status), (3, 'completed', 'completed'))
        self.client.post(visit_url, {'date': self.today.isoformat(), 'session_type': 'treatment'})
        pkg.refresh_from_db()
        self.assertEqual(pkg.sessions_used, 3)  # a completed package is not charged again
        self.assertEqual(Visit.objects.count(), 4)

    def test_attending_an_appointment_uses_a_session(self):
        self.package()
        appt = Appointment.objects.create(professional=self.profile, patient=self.patient, start=timezone.now(), end=timezone.now() + datetime.timedelta(minutes=30))
        self.client.post(reverse('accounts:appointment-action', args=[appt.pk]), {'action': 'attend'})
        self.assertEqual(PatientPackage.objects.get().sessions_used, 1)

    def test_cancelled_or_expired_packages_do_not_use_sessions(self):
        self.package()
        pkg = PatientPackage.objects.get()
        self.client.post(reverse('billing:package-cancel', args=[pkg.pk]))
        self.assertFalse(use_package_session(self.patient))
        pkg.status, pkg.start_date = 'active', self.today - datetime.timedelta(days=60)
        pkg.save()
        self.assertEqual(PatientPackage.objects.get().effective_status, 'expired')
        self.assertFalse(use_package_session(self.patient))

    def test_patient_billing_tab_shows_packages_and_takes_payments(self):
        self.package()
        pkg = PatientPackage.objects.get()
        tab = self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]), {'tab': 'billing'})
        self.assertContains(tab, '2 Weeks package')
        r = self.client.post(reverse('billing:package-payment', args=[pkg.pk]), {'amount': '500', 'payment_mode': 'online', 'next': 'patient'})
        self.assertRedirects(r, reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing')
        pkg.refresh_from_db()
        self.assertEqual(pkg.amount_paid, 1500)

    def test_packages_are_private(self):
        self.package()
        pkg = PatientPackage.objects.get()
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertNotContains(other.get(reverse('billing:packages')), self.patient.name)
        self.assertEqual(other.post(reverse('billing:package-payment', args=[pkg.pk]), {'amount': '1', 'payment_mode': 'cash'}).status_code, 404)
        self.assertEqual(other.post(reverse('billing:package-cancel', args=[pkg.pk])).status_code, 404)

    def test_revenue_combines_entries_and_package_payments(self):
        self.entry(rate='500')
        self.package()  # NPR 1000 paid at booking
        self.assertEqual(revenue_for(self.profile, self.today), Decimal('1500'))
        self.assertEqual(revenue_for(self.profile, self.today - datetime.timedelta(days=1)), 0)
        self.assertEqual(self.client.get(reverse('accounts:dashboard')).context['stats']['revenue_today'], Decimal('1500'))
        self.assertEqual(self.client.get(reverse('billing:billing')).context['package_payments'], Decimal('1000'))


class PatientBillTests(BillingBase):
    def test_bill_from_the_patient_tab_is_filled_from_the_record_and_returns_to_the_tab(self):
        tab = reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'
        page = self.client.get(tab)
        self.assertContains(page, reverse('billing:patient-bill', args=[self.patient.pk]))  # the small form, not a link to the whole book
        r = self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]),
                             {'entry_date': self.today.isoformat(), 'service': 'Ultrasound', 'rate': '700', 'payment_mode': 'online'})
        self.assertRedirects(r, tab, fetch_redirect_response=False)
        entry = BillingEntry.objects.get()
        self.assertEqual((entry.patient, entry.patient_name, entry.professional, entry.rate, entry.age), (self.patient, self.patient.name, self.profile, Decimal('700'), self.patient.age))
        self.assertContains(self.client.get(tab), 'Ultrasound')

    def test_negative_amount_is_refused_and_other_physios_patients_are_off_limits(self):
        url = reverse('billing:patient-bill', args=[self.patient.pk])
        self.client.post(url, {'entry_date': self.today.isoformat(), 'service': 'x', 'rate': '-5', 'payment_mode': 'cash'})
        self.assertEqual(BillingEntry.objects.count(), 0)
        other_user, other = make_professional('other')
        stranger = make_patient(other, name='Not Mine')
        self.assertEqual(self.client.post(reverse('billing:patient-bill', args=[stranger.pk]), {'entry_date': self.today.isoformat(), 'service': 'x', 'rate': '5', 'payment_mode': 'cash'}).status_code, 404)


class CreditBillTests(BillingBase):
    """Many patients pay at the end of the week, so a bill can be left on credit and settled later (in full or in parts)."""

    def test_a_bill_can_be_paid_in_full_part_paid_or_left_due(self):
        self.entry(patient_name='Full')
        self.entry(patient_name='Part', payment_status='partial', amount_received='200')
        self.entry(patient_name='Credit', payment_status='due', payment_mode='cash')
        full, part, due = (BillingEntry.objects.get(patient_name=n) for n in ('Full', 'Part', 'Credit'))
        self.assertEqual((full.status_key, full.amount_due), ('paid', 0))
        self.assertEqual((part.status_key, part.amount_paid, part.amount_due), ('partial', 200, 300))
        self.assertEqual((due.status_key, due.amount_paid, due.amount_due), ('due', 0, 500))
        page = self.client.get(reverse('billing:billing'))
        self.assertEqual(page.context['bills_due'], Decimal('800'))
        self.assertEqual(page.context['collected'], Decimal('700'))  # only what was received counts as money taken
        self.assertContains(page, 'Payments due')
        self.assertEqual(revenue_for(self.profile, self.today), Decimal('700'))

    def test_partial_needs_a_sensible_amount(self):
        self.entry(patient_name='Zero', payment_status='partial', amount_received='0')
        self.entry(patient_name='All', payment_status='partial', amount_received='500')
        self.assertEqual(BillingEntry.objects.count(), 0)

    def test_receiving_the_credit_later_counts_as_money_taken_on_that_day(self):
        yesterday = self.today - datetime.timedelta(days=3)
        self.entry(entry_date=yesterday.isoformat(), payment_status='due')
        entry = BillingEntry.objects.get()
        self.assertEqual(revenue_for(self.profile, yesterday), 0)
        r = self.client.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '300', 'payment_mode': 'online', 'next': '/billing/'})
        self.assertRedirects(r, '/billing/', fetch_redirect_response=False)
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, entry.status_key), (300, 'partial'))
        self.client.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '900', 'payment_mode': 'cash'})  # more than owed: capped
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, entry.status_key, entry.payments.count()), (500, 'paid', 2))
        self.assertEqual(revenue_for(self.profile, self.today), Decimal('500'))
        self.assertEqual(revenue_for(self.profile, yesterday), 0)
        self.assertEqual(self.client.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '10', 'payment_mode': 'cash'}).status_code, 302)
        self.assertEqual(entry.payments.count(), 2)  # already paid: nothing more recorded

    def test_credit_from_the_patient_tab_and_receive_there(self):
        tab = reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]),
                         {'entry_date': self.today.isoformat(), 'service': 'Therapy', 'rate': '900', 'payment_mode': 'cash', 'payment_status': 'due'})
        entry = BillingEntry.objects.get()
        self.assertEqual(entry.amount_due, 900)
        self.assertContains(self.client.get(tab), 'NPR 900.00 due')
        self.client.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '900', 'payment_mode': 'cash', 'next': tab})
        entry.refresh_from_db()
        self.assertEqual(entry.status_key, 'paid')

    def test_editing_cannot_drop_the_amount_below_what_was_received_and_strangers_are_refused(self):
        self.entry(payment_status='partial', amount_received='300')
        entry = BillingEntry.objects.get()
        self.client.post(reverse('billing:entry-edit', args=[entry.pk]), {'entry_date': self.today.isoformat(), 'patient_name': 'Walk In', 'service': 'Consult', 'rate': '100', 'payment_mode': 'cash'})
        entry.refresh_from_db()
        self.assertEqual(entry.rate, 500)
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '10', 'payment_mode': 'cash'}).status_code, 404)

    def test_existing_bills_were_all_paid(self):
        from importlib import import_module
        from django.apps import apps
        legacy = BillingEntry.objects.create(professional=self.profile, patient_name='Old', rate=400, payment_mode='card', entry_date=self.today)
        import_module('billing.migrations.0008_existing_bills_are_paid').mark_existing_paid(apps, None)
        legacy.refresh_from_db()
        self.assertEqual((legacy.amount_paid, legacy.payments.get().payment_mode), (400, 'card'))


class CancelBillTests(BillingBase):
    def test_a_paid_bill_can_still_be_edited_and_a_bill_can_be_cancelled_and_restored(self):
        self.entry(rate='500')
        entry = BillingEntry.objects.get()
        # edit after it was paid: raise the amount, the extra becomes due
        self.client.post(reverse('billing:entry-edit', args=[entry.pk]), {'entry_date': self.today.isoformat(), 'patient_name': 'Walk In', 'service': 'Consult + gel', 'rate': '650', 'payment_mode': 'cash'})
        entry.refresh_from_db()
        self.assertEqual((entry.rate, entry.amount_paid, entry.amount_due, entry.status_key), (650, 500, 150, 'partial'))
        self.assertContains(self.client.get(reverse('billing:billing'), {'edit': entry.pk}), f'/{entry.pk}/save/')
        # cancel: stays on record, counts for nothing
        r = self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'reason': 'Entered twice'})
        self.assertEqual(r.status_code, 302)
        entry.refresh_from_db()
        self.assertEqual((entry.status_key, entry.cancel_reason, entry.amount_due), ('cancelled', 'Entered twice', 0))
        page = self.client.get(reverse('billing:billing'))
        self.assertEqual((page.context['total'], page.context['bills_due'], page.context['collected'], page.context['month_total']), (0, 0, 0, 0))
        self.assertEqual(revenue_for(self.profile, self.today), 0)
        self.assertContains(page, 'Entered twice')
        self.assertIsNone(entry.record_payment(10))
        self.client.post(reverse('billing:entry-edit', args=[entry.pk]), {'entry_date': self.today.isoformat(), 'patient_name': 'Walk In', 'service': 'X', 'rate': '900', 'payment_mode': 'cash'})
        entry.refresh_from_db()
        self.assertEqual(entry.rate, 650)  # cancelled bills are not editable
        self.client.post(reverse('billing:entry-restore', args=[entry.pk]))
        entry.refresh_from_db()
        self.assertEqual((entry.status_key, revenue_for(self.profile, self.today)), ('partial', 500))

    def test_cancel_from_the_patient_tab_and_ownership(self):
        tab = reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'T', 'rate': '300', 'payment_mode': 'cash', 'payment_status': 'due'})
        entry = BillingEntry.objects.get()
        self.assertContains(self.client.get(tab), 'Cancel bill')
        self.assertRedirects(self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'next': tab}), tab, fetch_redirect_response=False)
        self.assertContains(self.client.get(tab), 'Cancelled')
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.post(reverse('billing:entry-cancel', args=[entry.pk])).status_code, 404)


class CorrectBillTests(BillingBase):
    def make_bill(self, **extra):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Therapy', 'rate': '800', 'payment_mode': 'cash', **extra})
        return BillingEntry.objects.get()

    def correct(self, entry, **extra):
        data = {'entry_date': self.today.isoformat(), 'service': 'Therapy', 'rate': '800', 'payment_mode': 'cash', 'payment_status': 'paid', **extra}
        return self.client.post(reverse('billing:entry-correct', args=[entry.pk]), data)

    def test_reception_can_turn_a_paid_bill_into_credit_and_back_without_touching_patient_details(self):
        entry = self.make_bill()  # recorded as paid in full by mistake
        self.assertEqual(entry.status_key, 'paid')
        before = (entry.patient_name, entry.age, entry.sex, entry.contact_number)
        self.correct(entry, payment_status='due')
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, entry.amount_due, entry.payments.count()), (0, 800, 0))
        self.assertEqual(revenue_for(self.profile, self.today), 0)
        self.assertEqual((entry.patient_name, entry.age, entry.sex, entry.contact_number), before)
        self.correct(entry, payment_status='partial', amount_received='300', payment_mode='online')
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, entry.status_key, entry.payments.get().payment_mode), (300, 'partial', 'online'))
        self.correct(entry, payment_status='paid', rate='700')  # wrong amount too
        entry.refresh_from_db()
        self.assertEqual((entry.rate, entry.amount_paid, entry.status_key, sum(p.amount for p in entry.payments.all())), (700, 700, 'paid', 700))
        self.assertEqual(revenue_for(self.profile, self.today), 700)

    def test_lowering_what_was_received_shrinks_the_latest_payments(self):
        entry = self.make_bill(payment_status='partial', amount_received='500')
        entry.record_payment(300, 'cash')
        self.correct(entry, payment_status='partial', amount_received='400')
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, sum(p.amount for p in entry.payments.all())), (400, 400))

    def test_validation_cancelled_and_ownership(self):
        entry = self.make_bill()
        self.correct(entry, payment_status='partial', amount_received='800')  # not partial
        self.correct(entry, rate='-5')
        entry.refresh_from_db()
        self.assertEqual((entry.rate, entry.amount_paid), (800, 800))
        self.client.post(reverse('billing:entry-cancel', args=[entry.pk]))
        self.correct(entry, rate='100')
        entry.refresh_from_db()
        self.assertEqual(entry.rate, 800)
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.post(reverse('billing:entry-correct', args=[entry.pk]), {'rate': '1'}).status_code, 404)

    def test_the_patient_tab_and_the_book_open_the_small_edit_form(self):
        entry = self.make_bill(payment_status='due')
        tab = self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing')
        self.assertContains(tab, reverse('billing:entry-correct', args=[entry.pk]))
        self.assertContains(self.client.get(reverse('billing:billing'), {'edit': entry.pk}), reverse('billing:entry-correct', args=[entry.pk]))


class BillingSearchTests(BillingBase):
    def test_searching_a_patient_shows_only_that_patients_bills_across_all_dates(self):
        other = make_patient(self.profile, 'Other Person')
        last_week = self.today - datetime.timedelta(days=6)
        for pat, day, rate, extra in ((self.patient, self.today, '400', {}), (self.patient, last_week, '600', {'payment_status': 'due'}), (other, self.today, '999', {})):
            self.client.post(reverse('billing:patient-bill', args=[pat.pk]), {'entry_date': day.isoformat(), 'service': 'Therapy', 'rate': rate, 'payment_mode': 'cash', **extra})
        page = self.client.get(reverse('billing:billing'), {'patient': self.patient.pk})
        self.assertEqual([e.rate for e in page.context['entries']], [400, 600])  # newest first, both dates, nobody else
        s = page.context['search']
        self.assertEqual((s['count'], s['billed'], s['received'], s['due']), (2, 1000, 400, 600))
        self.assertEqual([e.pk for e in page.context['owed']], [e.pk for e in page.context['entries'] if e.amount_due])
        self.assertNotIn('Other Person', [e.patient_name for e in page.context['entries']])
        self.assertContains(page, 'Clear')
        # the normal day view is untouched
        self.assertIsNone(self.client.get(reverse('billing:billing')).context['search'])

    def test_text_search_and_privacy(self):
        self.entry(patient_name='Walk In Sita', contact_number='9800000123')
        self.entry(patient_name='Someone Else')
        page = self.client.get(reverse('billing:billing'), {'q': '9800000123'})
        self.assertEqual([e.patient_name for e in page.context['entries']], ['Walk In Sita'])
        self.assertEqual(len(self.client.get(reverse('billing:billing'), {'q': 'nobody'}).context['entries']), 0)
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(len(other.get(reverse('billing:billing'), {'patient': self.patient.pk, 'q': 'Walk'}).context['entries']), 0)


class AdvanceTests(BillingBase):
    """A patient pays NPR 1000 ahead; each session of 700 is taken from it."""

    def take(self, amount='1000', **extra):
        return self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': amount, 'payment_mode': 'cash', **extra})

    def bill(self, rate='700', **extra):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Session', 'rate': rate, 'payment_mode': 'cash', 'payment_status': 'advance', **extra})

    def test_advance_is_money_taken_once_and_sessions_draw_it_down(self):
        self.take('1000')
        self.assertEqual(advance_balance(self.patient.pk), 1000)
        self.assertEqual(revenue_for(self.profile, self.today), 1000)  # the cash arrived today
        self.bill('700')
        first = BillingEntry.objects.get()
        self.assertEqual((first.status_key, first.amount_paid, advance_balance(self.patient.pk)), ('paid', 700, 300))
        self.assertEqual(revenue_for(self.profile, self.today), 1000)  # settling from the advance is not new money
        self.bill('700')  # only 300 left: the rest stays due
        second = BillingEntry.objects.exclude(pk=first.pk).get()
        self.assertEqual((second.amount_paid, second.amount_due, second.status_key, advance_balance(self.patient.pk)), (300, 400, 'partial', 0))
        page = self.client.get(reverse('billing:billing'))
        self.assertEqual((page.context['received_total'], page.context['collected'], page.context['advance_net'], page.context['advance_held']), (1000, 0, 1000, 0))
        self.assertEqual(dict(page.context['by_mode'])['Cash'], 1000)
        tab = self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing')
        self.assertContains(tab, 'Advance balance: NPR 0')
        self.assertContains(tab, 'Used for a bill')

    def test_take_advance_can_clear_unpaid_bills_and_credit_can_be_received_from_the_advance(self):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '600', 'payment_mode': 'cash', 'payment_status': 'due'})
        self.take('1000', settle='1')
        entry = BillingEntry.objects.get()
        self.assertEqual((entry.status_key, advance_balance(self.patient.pk)), ('paid', 400))
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '900', 'payment_mode': 'cash', 'payment_status': 'due'})
        late = BillingEntry.objects.exclude(pk=entry.pk).get()
        self.client.post(reverse('billing:entry-payment', args=[late.pk]), {'amount': '250', 'payment_mode': 'advance'})
        late.refresh_from_db()
        self.assertEqual((late.amount_paid, advance_balance(self.patient.pk)), (250, 150))
        self.client.post(reverse('billing:entry-payment', args=[late.pk]), {'amount': '900', 'payment_mode': 'advance'})  # capped by what is held
        late.refresh_from_db()
        self.assertEqual((late.amount_paid, advance_balance(self.patient.pk)), (400, 0))

    def test_cancelling_a_bill_returns_its_advance_and_refunds_are_money_out(self):
        self.take('1000')
        self.bill('700')
        entry = BillingEntry.objects.get()
        self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'reason': 'x'})
        self.assertEqual(advance_balance(self.patient.pk), 1000)
        r = self.client.post(reverse('billing:advance-refund', args=[self.patient.pk]), {'amount': '400', 'payment_mode': 'cash'})
        self.assertEqual((r.status_code, advance_balance(self.patient.pk)), (302, 600))
        self.assertEqual(revenue_for(self.profile, self.today), 600)  # 1000 taken less 400 paid back
        self.client.post(reverse('billing:advance-refund', args=[self.patient.pk]), {'amount': '9999', 'payment_mode': 'cash'})
        self.assertEqual(advance_balance(self.patient.pk), 600)  # cannot refund more than is held

    def test_correcting_a_bill_that_used_the_advance_puts_it_back_and_the_rules_hold(self):
        self.take('1000')
        self.bill('700')
        entry = BillingEntry.objects.get()
        self.client.post(reverse('billing:entry-correct', args=[entry.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '700', 'payment_mode': 'cash', 'payment_status': 'due'})
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, advance_balance(self.patient.pk)), (0, 1000))
        self.assertEqual(self.take('0').status_code, 302)
        self.assertEqual(advance_balance(self.patient.pk), 1000)  # zero is refused
        self.bill('50')  # no balance question here: balance is 1000, fine
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '5', 'payment_mode': 'cash'}).status_code, 404)

    def test_no_advance_leaves_the_bill_on_credit_and_walk_ins_cannot_use_one(self):
        self.bill('700')
        self.assertEqual(BillingEntry.objects.get().status_key, 'due')
        self.entry(patient_name='Walk', payment_status='advance')
        self.assertEqual(BillingEntry.objects.get(patient_name='Walk').amount_paid, 0)


class CarryOverTests(BillingBase):
    """The patient hands over NPR 1000 for a 700 session and there is no change: 300 is carried over to tomorrow."""

    def bill(self, rate='700', **extra):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Session', 'rate': rate, 'payment_mode': 'cash', **extra})

    def test_extra_cash_with_no_change_is_carried_over_and_settles_tomorrows_bill(self):
        self.bill('700', payment_status='paid', amount_given='1000')
        first = BillingEntry.objects.get()
        self.assertEqual((first.status_key, first.amount_paid, advance_balance(self.patient.pk)), ('paid', 700, 300))
        self.assertEqual(revenue_for(self.profile, self.today), 1000)  # all 1000 was handed over today
        # tomorrow: 300 from the advance, the patient pays only the other 400
        self.bill('700', payment_status='advance_cash', amount_given='400')
        second = BillingEntry.objects.exclude(pk=first.pk).get()
        self.assertEqual((second.status_key, second.amount_paid, advance_balance(self.patient.pk)), ('paid', 700, 0))
        self.assertEqual(sorted(p.payment_mode for p in second.payments.all()), ['advance', 'cash'])
        self.assertEqual(revenue_for(self.profile, self.today), 1400)
        self.assertContains(self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'), 'Change not returned')

    def test_blank_cash_means_exact_and_a_larger_second_payment_carries_again(self):
        self.bill('700', payment_status='paid')
        self.assertEqual(advance_balance(self.patient.pk), 0)
        self.bill('700', payment_status='paid', amount_given='1000')
        self.bill('700', payment_status='advance_cash', amount_given='500')  # 300 advance + 400 owed, 500 given: 100 carried
        self.assertEqual(advance_balance(self.patient.pk), 100)
        self.bill('700', payment_status='advance_cash')  # blank given pays exactly the rest (600) in cash after using the 100
        self.assertEqual(advance_balance(self.patient.pk), 0)
        self.assertTrue(all(e.status_key == 'paid' for e in BillingEntry.objects.all()))

    def test_validation_and_walk_ins(self):
        self.bill('700', payment_status='paid', amount_given='500')  # less than the bill
        self.assertEqual(BillingEntry.objects.count(), 0)
        self.entry(patient_name='Walk', rate='700', amount_given='1000')  # walk-in cannot carry over
        self.assertEqual(BillingEntry.objects.count(), 0)
        self.entry(patient_name='Walk', rate='700', amount_given='700')
        self.assertEqual(BillingEntry.objects.get().status_key, 'paid')

    def test_receiving_a_due_bill_with_more_cash_carries_the_extra(self):
        self.bill('400', payment_status='due')
        entry = BillingEntry.objects.get()
        self.client.post(reverse('billing:entry-payment', args=[entry.pk]), {'amount': '1000', 'payment_mode': 'cash'})
        entry.refresh_from_db()
        self.assertEqual((entry.status_key, advance_balance(self.patient.pk)), ('paid', 600))
        self.assertEqual(revenue_for(self.profile, self.today), 1000)


class AdvanceTabTests(BillingBase):
    def test_the_advances_tab_lists_holders_and_shows_one_patients_full_history(self):
        other = make_patient(self.profile, 'Maya Rai')
        self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '1000', 'payment_mode': 'cash'})
        self.client.post(reverse('billing:advance-add', args=[other.pk]), {'amount': '500', 'payment_mode': 'online'})
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '700', 'payment_mode': 'cash', 'payment_status': 'advance'})
        self.client.post(reverse('billing:advance-refund', args=[other.pk]), {'amount': '500', 'payment_mode': 'cash'})  # Maya is now at zero
        page = self.client.get(reverse('billing:advances'))
        self.assertEqual((page.context['held'], page.context['holders']), (300, 1))
        self.assertEqual([r['patient'].pk for r in page.context['rows']], [self.patient.pk])  # only people with a balance by default
        chosen = page.context['chosen']
        self.assertEqual((chosen['balance'], chosen['taken'], chosen['used'], chosen['refunded']), (300, 1000, 700, 0))
        self.assertEqual([r['running'] for r in page.context['ledger']], [300, 1000])  # newest first, with the balance after each
        self.assertContains(page, 'Used for a bill')
        everyone = self.client.get(reverse('billing:advances'), {'show': 'all'})
        self.assertEqual(len(everyone.context['rows']), 2)
        detail = self.client.get(reverse('billing:advances'), {'patient': other.pk, 'show': 'all'})
        self.assertEqual((detail.context['chosen']['balance'], detail.context['chosen']['refunded']), (0, 500))
        self.assertEqual(len(self.client.get(reverse('billing:advances'), {'q': 'zzz', 'show': 'all'}).context['rows']), 0)
        self.assertContains(self.client.get(reverse('billing:billing')), reverse('billing:advances'))  # the card opens the tab

    def test_actions_from_the_tab_return_to_it_and_other_professionals_see_nothing(self):
        url = reverse('billing:advances') + f'?patient={self.patient.pk}'
        r = self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '200', 'payment_mode': 'cash', 'next': url})
        self.assertRedirects(r, url, fetch_redirect_response=False)
        r = self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '50', 'payment_mode': 'cash', 'next': 'https://evil.example/'})
        self.assertRedirects(r, reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing', fetch_redirect_response=False)
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        page = other.get(reverse('billing:advances'), {'patient': self.patient.pk})
        self.assertEqual((page.context['held'], page.context['rows'], page.context['chosen']), (0, [], None))

    def test_links_from_billing_open_the_patients_billing_tab_directly(self):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '700', 'payment_mode': 'cash'})
        self.package()
        target = reverse('accounts:patient-detail', args=[self.patient.pk])
        self.assertContains(self.client.get(reverse('billing:billing')), f'{target}?tab=billing')
        self.assertNotContains(self.client.get(reverse('billing:billing')), f'href="{target}"')
        self.assertContains(self.client.get(reverse('billing:packages')), f'{target}?tab=billing')


class InvoiceTests(BillingBase):
    def make(self, **extra):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Therapy session', 'rate': '700', 'payment_mode': 'cash', 'payment_status': 'paid', **extra})
        return BillingEntry.objects.latest('id')

    def test_the_invoice_page_offers_pdf_whatsapp_email_and_copy(self):
        entry = self.make()
        page = self.client.get(reverse('billing:invoice', args=[entry.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, entry.invoice_number)
        self.assertContains(page, reverse('billing:invoice-pdf', args=[entry.pk]))
        wa = page.context['whatsapp']
        self.assertTrue(wa.startswith('https://wa.me/9779841000000?text='), wa)  # a local number gets the country code
        self.assertIn(entry.invoice_number, page.context['message'])
        self.assertContains(page, 'Send by e-mail')
        self.assertContains(self.client.get(reverse('billing:billing')), reverse('billing:invoice', args=[entry.pk]))
        self.assertContains(self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'), reverse('billing:invoice', args=[entry.pk]))

    def test_pdf_is_a_real_pdf_for_the_owner_only(self):
        entry = self.make(payment_status='partial', amount_received='200')
        r = self.client.get(reverse('billing:invoice-pdf', args=[entry.pk]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/pdf'))
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('attachment', self.client.get(reverse('billing:invoice-pdf', args=[entry.pk]), {'download': '1'})['Content-Disposition'])
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.get(reverse('billing:invoice-pdf', args=[entry.pk])).status_code, 404)
        self.assertEqual(other.get(reverse('billing:invoice', args=[entry.pk])).status_code, 404)

    def test_the_shared_link_works_without_signing_in_and_only_for_that_bill(self):
        from . import invoice as inv
        entry = self.make()
        token = inv.invoice_token(entry)
        anonymous = self.client_class()
        r = anonymous.get(reverse('billing:invoice-public', args=[token]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/pdf'))
        self.assertEqual(anonymous.get(reverse('billing:invoice-public', args=['not-a-real-token'])).status_code, 404)
        self.assertEqual(anonymous.get(reverse('billing:invoice-public', args=[token[:-3] + 'abc'])).status_code, 404)
        from django.core import signing
        self.assertIsNone(inv.entry_from_token(signing.dumps(entry.pk, salt='something-else')))
        entry.delete()
        self.assertEqual(anonymous.get(reverse('billing:invoice-public', args=[token])).status_code, 404)

    def test_email_sends_the_pdf_attached_and_rejects_bad_addresses(self):
        from django.core import mail
        entry = self.make()
        self.client.post(reverse('billing:invoice-email', args=[entry.pk]), {'email': 'not-an-email'})
        self.assertEqual(len(mail.outbox), 0)
        self.client.post(reverse('billing:invoice-email', args=[entry.pk]), {'email': 'sita@example.com'})
        self.assertEqual(len(mail.outbox), 1)
        sent = mail.outbox[0]
        self.assertEqual(sent.to, ['sita@example.com'])
        self.assertIn(entry.invoice_number, sent.subject)
        name, content, kind = sent.attachments[0]
        self.assertEqual((name, kind), (f'{entry.invoice_number}.pdf', 'application/pdf'))
        self.assertTrue(content.startswith(b'%PDF'))
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        other.post(reverse('billing:invoice-email', args=[entry.pk]), {'email': 'x@example.com'})
        self.assertEqual(len(mail.outbox), 1)

    def test_phone_numbers_non_latin_text_and_cancelled_bills(self):
        from . import invoice as inv
        self.assertEqual(inv.normalize_phone('9841234567'), '9779841234567')
        self.assertEqual(inv.normalize_phone('+977 984-1234567'), '9779841234567')
        self.assertEqual(inv.normalize_phone('09841234567'), '9779841234567')
        self.assertEqual(inv.normalize_phone('0097798412345'), '97798412345')
        self.assertEqual(inv.normalize_phone('12'), '')
        entry = self.make(service='पिठ्यूँको उपचार')
        entry.patient_name = 'सीता गुरुङ'
        entry.save()
        self.assertTrue(inv.build_pdf(entry).startswith(b'%PDF'))  # a Nepali name or service never breaks the PDF
        self.assertContains(self.client.get(reverse('billing:invoice', args=[entry.pk])), 'सीता गुरुङ')  # the online invoice shows it in full
        self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'reason': 'x'})
        self.assertContains(self.client.get(reverse('billing:invoice', args=[entry.pk])), 'CANCELLED')
        entry.refresh_from_db()
        self.assertTrue(inv.build_pdf(entry).startswith(b'%PDF'))


class InvoiceExtraTests(BillingBase):
    def test_the_invoice_states_the_extra_kept_as_advance(self):
        from . import invoice as inv
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Session', 'rate': '500', 'payment_mode': 'cash', 'payment_status': 'paid', 'amount_given': '700'})
        entry = BillingEntry.objects.get()
        self.assertEqual((entry.rate, entry.amount_paid, entry.carried_over, entry.cash_handed_over), (500, 500, 200, 700))
        page = self.client.get(reverse('billing:invoice', args=[entry.pk]))
        self.assertContains(page, 'Extra kept as advance: NPR 200')
        self.assertContains(page, 'You handed over NPR 700.00 for a NPR 500.00 bill')
        self.assertEqual((page.context['carried'], page.context['advance_now']), (200, 200))
        self.assertIn('Your extra NPR 200.00 is kept as advance', page.context['message'])
        self.assertTrue(self.client.get(reverse('billing:invoice-pdf', args=[entry.pk])).content.startswith(b'%PDF'))
        self.assertContains(self.client.get(reverse('billing:billing')), '+NPR 200')
        self.assertContains(self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'), '+NPR 200')

    def test_a_bill_with_no_extra_says_nothing_about_advance(self):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Session', 'rate': '500', 'payment_mode': 'cash', 'payment_status': 'paid'})
        entry = BillingEntry.objects.get()
        page = self.client.get(reverse('billing:invoice', args=[entry.pk]))
        self.assertNotContains(page, 'Extra kept as advance')
        self.assertNotIn('extra', page.context['message'])


class InvoicePictureTests(BillingBase):
    def test_the_invoice_page_can_send_a_picture_without_a_link(self):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': '500', 'payment_mode': 'cash', 'payment_status': 'paid'})
        entry = BillingEntry.objects.get()
        page = self.client.get(reverse('billing:invoice', args=[entry.pk]))
        for needle in ('id="picShare"', 'id="picWa"', 'id="picSave"', 'vendor/html2canvas/html2canvas.min.js'):
            self.assertContains(page, needle)
        self.assertEqual(page.context['whatsapp_chat'], 'https://wa.me/9779841000000')  # the chat only, no message and no link
        from . import invoice as inv
        self.assertEqual(inv.whatsapp_chat(''), 'https://web.whatsapp.com/')
        from django.contrib.staticfiles import finders
        self.assertTrue(finders.find('vendor/html2canvas/html2canvas.min.js'))


class InvoicePngTests(BillingBase):
    def test_the_server_draws_the_invoice_picture_as_a_fallback(self):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'Session', 'rate': '500', 'payment_mode': 'cash', 'payment_status': 'paid', 'amount_given': '700'})
        entry = BillingEntry.objects.get()
        r = self.client.get(reverse('billing:invoice-png', args=[entry.pk]))
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'image/png'))
        self.assertTrue(r.content.startswith(b'\x89PNG'))
        entry.patient_name, entry.service = 'सीता गुरुङ', 'पिठ्यूँको उपचार'  # non-Latin text never breaks it
        entry.save()
        self.assertTrue(self.client.get(reverse('billing:invoice-png', args=[entry.pk])).content.startswith(b'\x89PNG'))
        self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'reason': 'x'})
        self.assertEqual(self.client.get(reverse('billing:invoice-png', args=[entry.pk])).status_code, 200)
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.get(reverse('billing:invoice-png', args=[entry.pk])).status_code, 404)
        self.assertContains(self.client.get(reverse('billing:invoice', args=[entry.pk])), reverse('billing:invoice-png', args=[entry.pk]))


class PackageDaysTests(BillingBase):
    def start(self, **extra):
        return self.client.post(reverse('billing:package-create'), {'patient': self.patient.pk, 'total_sessions': 10, 'total_amount': '5000', 'amount_paid_now': '0', 'start_date': self.today.isoformat(), **extra})

    def test_ten_day_package_from_the_list_and_any_number_of_days_from_the_box(self):
        self.start(duration='10_day')
        ten = PatientPackage.objects.get()
        self.assertEqual((ten.duration, ten.length_days, ten.get_duration_display(), ten.end_date), ('10_day', 10, '10 Days', self.today + datetime.timedelta(days=10)))
        self.start(duration='1_week', custom_days='12')  # typed days win over the list
        twelve = PatientPackage.objects.exclude(pk=ten.pk).get()
        self.assertEqual((twelve.duration, twelve.custom_days, twelve.length_days, twelve.get_duration_display()), ('custom', 12, 12, '12 Days'))
        self.assertEqual(twelve.end_date, self.today + datetime.timedelta(days=12))
        self.start(duration='', custom_days='1')
        self.assertEqual(PatientPackage.objects.exclude(pk__in=[ten.pk, twelve.pk]).get().get_duration_display(), '1 Day')
        page = self.client.get(reverse('billing:packages'))
        self.assertContains(page, '10 Days package')
        self.assertContains(page, '12 Days package')
        self.assertContains(page, 'Or type the number of days')
        self.assertContains(page, '10 Days</option>')
        self.assertNotContains(page, 'Other (number of days)')  # the box is the way to type it, not a list entry
        self.assertContains(self.client.get(reverse('accounts:patient-detail', args=[self.patient.pk]) + '?tab=billing'), '12 Days package')

    def test_neither_a_choice_nor_days_or_a_silly_number_is_refused(self):
        self.start(duration='')
        self.start(duration='1_week', custom_days='0')
        self.start(duration='1_week', custom_days='9999')
        self.assertEqual(PatientPackage.objects.count(), 0)
        self.start(duration='1_week')  # the old way still works
        self.assertEqual(PatientPackage.objects.get().length_days, 7)


class MoneySafetyTests(BillingBase):
    def bill(self, rate='500', **extra):
        self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), {'entry_date': self.today.isoformat(), 'service': 'S', 'rate': rate, 'payment_mode': 'cash', 'payment_status': 'due', **extra})
        return BillingEntry.objects.latest('id')

    def test_two_people_receiving_the_same_bill_cannot_both_count_the_balance(self):
        entry = self.bill()
        first, second = BillingEntry.objects.get(pk=entry.pk), BillingEntry.objects.get(pk=entry.pk)  # two screens open on the same unpaid bill
        self.assertIsNotNone(first.record_payment(500, 'cash'))
        self.assertIsNone(second.record_payment(500, 'cash'))  # the second one sees it is already paid
        entry.refresh_from_db()
        self.assertEqual((entry.amount_paid, entry.payments.count()), (500, 1))
        part = self.bill()
        a, b = BillingEntry.objects.get(pk=part.pk), BillingEntry.objects.get(pk=part.pk)
        a.record_payment(300, 'cash')
        paid = b.record_payment(300, 'cash')  # only 200 was still owed
        self.assertEqual(paid.amount, 200)
        part.refresh_from_db()
        self.assertEqual((part.amount_paid, part.status_key), (500, 'paid'))

    def test_the_same_advance_cannot_be_spent_twice(self):
        self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '500', 'payment_mode': 'cash'})
        one, two = self.bill(), self.bill()
        a, b = BillingEntry.objects.get(pk=one.pk), BillingEntry.objects.get(pk=two.pk)
        self.assertEqual(a.use_advance(), 500)
        self.assertEqual(b.use_advance(), 0)  # nothing left
        self.assertEqual(advance_balance(self.patient.pk), 0)

    def test_a_double_refund_is_refused(self):
        self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '500', 'payment_mode': 'cash'})
        for _ in range(2):
            self.client.post(reverse('billing:advance-refund', args=[self.patient.pk]), {'amount': '500', 'payment_mode': 'cash'})
        self.assertEqual(advance_balance(self.patient.pk), 0)  # not -500

    def test_package_payments_use_the_latest_balance_too(self):
        self.client.post(reverse('billing:package-create'), {'patient': self.patient.pk, 'duration': '10_day', 'total_sessions': 5, 'total_amount': '1000', 'amount_paid_now': '0', 'start_date': self.today.isoformat()})
        package = PatientPackage.objects.get()
        a, b = PatientPackage.objects.get(pk=package.pk), PatientPackage.objects.get(pk=package.pk)
        a.record_payment(800)
        self.assertEqual(b.record_payment(800).amount, 200)
        package.refresh_from_db()
        self.assertEqual(package.amount_paid, 1000)

    def test_check_money_finds_drift_and_is_quiet_when_all_is_well(self):
        from io import StringIO

        from django.core.management import call_command
        from django.core.management.base import CommandError
        entry = self.bill()
        entry.record_payment(200, 'cash')
        self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '300', 'payment_mode': 'cash'})
        out = StringIO()
        call_command('check_money', stdout=out)
        self.assertIn('Money checks out', out.getvalue())
        BillingEntry.objects.filter(pk=entry.pk).update(amount_paid=450)  # a figure edited by hand, away from its payments
        with self.assertRaises(CommandError):
            call_command('check_money', stdout=out, stderr=StringIO())
        self.assertIn(entry.invoice_number, out.getvalue())


class BillingSheetTests(BillingBase):
    def bill(self, name=None, rate='500', status='paid', day=None, **extra):
        data = {'entry_date': (day or self.today).isoformat(), 'service': 'Therapy', 'rate': rate, 'payment_mode': 'cash', 'payment_status': status, **extra}
        if name:
            self.client.post(reverse('billing:entry-create'), {**data, 'patient_name': name})
        else:
            self.client.post(reverse('billing:patient-bill', args=[self.patient.pk]), data)
        return BillingEntry.objects.latest('id')

    def workbook(self, **params):
        import io

        from openpyxl import load_workbook
        r = self.client.get(reverse('billing:export-xlsx'), params)
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        self.assertIn('attachment', r['Content-Disposition'])
        return load_workbook(io.BytesIO(r.content))

    def test_the_sheet_shows_bills_with_totals_and_filters(self):
        self.bill(rate='500')
        self.bill(rate='700', status='due')
        self.bill(name='Walk In Ram', rate='300', status='partial', amount_received='100')
        old = self.bill(rate='900', day=self.today - datetime.timedelta(days=60))
        page = self.client.get(reverse('billing:sheet'))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context['totals']['count'], 3)  # this month only by default: the old bill is outside
        self.assertEqual((page.context['totals']['amount'], page.context['totals']['received'], page.context['totals']['due']), (1500, 600, 900))
        self.assertContains(page, 'Download Excel')
        self.assertEqual(len(self.client.get(reverse('billing:sheet'), {'from': (self.today - datetime.timedelta(days=90)).isoformat()}).context['rows']), 4)
        self.assertEqual(len(self.client.get(reverse('billing:sheet'), {'status': 'due'}).context['rows']), 1)
        self.assertEqual(len(self.client.get(reverse('billing:sheet'), {'status': 'unpaid'}).context['rows']), 2)
        self.assertEqual([r['patient'] for r in self.client.get(reverse('billing:sheet'), {'q': 'ram'}).context['rows']], ['Walk In Ram'])
        self.assertEqual(len(self.client.get(reverse('billing:sheet'), {'from': 'junk', 'to': 'junk', 'status': 'bogus', 'mode': 'bogus'}).context['rows']), 3)  # bad input falls back to the defaults
        self.assertContains(self.client.get(reverse('billing:billing')), reverse('billing:sheet'))  # the tab

    def test_the_excel_file_has_the_sheets_the_rows_and_the_summary(self):
        self.bill(rate='500')
        self.bill(rate='700', status='due')
        self.client.post(reverse('billing:advance-add', args=[self.patient.pk]), {'amount': '1000', 'payment_mode': 'online'})
        wb = self.workbook()
        self.assertEqual(wb.sheetnames, ['Summary', 'Bills', 'Payments received', 'Advances', 'Packages'])
        bills = wb['Bills']
        self.assertEqual([c.value for c in bills[1]][:9], ['Date', 'Invoice', 'Patient', 'Patient ID', 'Phone', 'Service', 'Amount (NPR)', 'Received (NPR)', 'Due (NPR)'])
        self.assertEqual((bills['A2'].value.date() if hasattr(bills['A2'].value, 'date') else bills['A2'].value), self.today)
        self.assertEqual([bills.cell(row=r, column=7).value for r in (2, 3)], [500, 700])
        self.assertEqual(bills.cell(row=3, column=9).value, 700)
        self.assertEqual(bills.cell(row=4, column=7).value, '=SUBTOTAL(109,G2:G3)')  # a total that follows the filter you apply inside Excel
        self.assertEqual(bills.freeze_panes, 'A2')
        summary = {row[0].value: row[1].value for row in wb['Summary'].iter_rows(min_row=4, max_row=8)}
        self.assertEqual((summary['Bills'], summary['Billed (NPR)'], summary['Received on these bills (NPR)']), (2, 1200, 500))
        flat = [c.value for row in wb['Summary'].iter_rows() for c in row]
        self.assertIn('Total received', flat)
        self.assertEqual(len(list(wb['Payments received'].iter_rows(min_row=2))), 1)
        self.assertEqual([c.value for c in wb['Advances'][2]][3:5], ['Advance taken', 1000])
        self.assertEqual(self.client.get(reverse('billing:export-xlsx'), {'status': 'due'}).status_code, 200)

    def test_text_that_looks_like_a_formula_stays_text_in_excel_and_csv(self):
        self.bill(name='=HYPERLINK("http://evil","x")', rate='100')
        self.bill(name='+cmd|calc', rate='100')
        wb = self.workbook()
        cell = wb['Bills'].cell(row=2, column=3)
        self.assertEqual(cell.data_type, 's')  # stored as text, never as a formula
        r = self.client.get(reverse('billing:export-csv'))
        text = r.content.decode('utf-8-sig')
        self.assertEqual(r['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn("'=HYPERLINK", text)
        self.assertIn("'+cmd|calc", text)
        self.assertTrue(r.content.startswith(b'\xef\xbb\xbf'))  # the byte-order mark Excel needs for Nepali text
        self.assertIn('Invoice', text.splitlines()[0])

    def test_nepali_names_and_cancelled_bills_export(self):
        entry = self.bill(name='सीता गुरुङ', rate='500')
        self.client.post(reverse('billing:entry-cancel', args=[entry.pk]), {'reason': 'x'})
        self.bill(rate='300')
        wb = self.workbook()
        names = [wb['Bills'].cell(row=r, column=3).value for r in (2, 3)]
        self.assertIn('सीता गुरुङ', names)
        statuses = [wb['Bills'].cell(row=r, column=10).value for r in (2, 3)]
        self.assertIn('Cancelled', statuses)
        summary = {row[0].value: row[1].value for row in wb['Summary'].iter_rows(min_row=4, max_row=8)}
        self.assertEqual(summary['Billed (NPR)'], 300)  # a cancelled bill is listed but not counted
        page = self.client.get(reverse('billing:sheet'))
        self.assertContains(page, 'सीता गुरुङ')
        self.assertEqual(page.context['totals']['amount'], 300)

    def test_the_sheet_and_downloads_are_private_and_need_an_approved_workspace(self):
        self.bill(rate='500')
        stranger, _ = make_professional('stranger')
        other = self.client_class()
        other.force_login(stranger)
        self.assertEqual(other.get(reverse('billing:sheet')).context['totals']['count'], 0)
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(other.get(reverse('billing:export-xlsx')).content))
        self.assertEqual(wb['Bills'].max_row, 1)  # headers only: another professional's bills are never in it
        make_professional('waiting', status='submitted')
        waiting = self.client_class()
        waiting.login(username='waiting', password='Gr33n!Health-99')
        self.assertEqual(waiting.get(reverse('billing:export-xlsx')).status_code, 302)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('billing:sheet')).status_code, 302)
        self.assertEqual(self.client.get(reverse('billing:export-csv')).status_code, 302)

    def test_a_long_sheet_is_cut_in_the_browser_but_not_in_the_download(self):
        from . import exports
        old = exports.MAX_SHEET_ROWS
        exports.MAX_SHEET_ROWS = 2
        try:
            for _ in range(3):
                self.bill(rate='100')
            page = self.client.get(reverse('billing:sheet'))
            self.assertTrue(page.context['truncated'])
            self.assertEqual(len(page.context['rows']), 2)
            self.assertContains(page, 'Showing the first 2 rows')
        finally:
            exports.MAX_SHEET_ROWS = old
        self.assertEqual(self.workbook()['Bills'].max_row, 1 + 3 + 1)  # header, three bills and the total row
